from pathlib import Path
import csv, json
import pymupdf as fitz
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'livrables'; OUT.mkdir(exist_ok=True)
font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
pdfmetrics.registerFont(TTFont('Institution',font))
pdfmetrics.registerFont(TTFont('InstitutionBold',font.replace('.ttf','-Bold.ttf')))
W,H=1190.55,841.89
navy='#102D49'; blue='#25618A'; pale='#EDF5FA'; gray='#F1F4F6'; ink='#17364D'; line='#8095A5'
prs=Presentation(); prs.slide_width=Pt(W); prs.slide_height=Pt(H)
slide=prs.slides.add_slide(prs.slide_layouts[6])
c=canvas.Canvas(str(OUT/'Organigramme_modernise_A3.pdf'),pagesize=(W,H)); c.setTitle('Organigramme institutionnel — version modernisée')
nodes=[]; edges=[]
def rgb(s):return RGBColor.from_string(s.lstrip('#'))
def text(x,y,w,h,txt,size=12,color=ink,bold=False,align='left'):
    lines=txt.split('\n'); leading=size*1.3
    c.setFillColor(HexColor(color)); c.setFont('InstitutionBold' if bold else 'Institution',size)
    yy=H-y-(h-len(lines)*leading)/2-size
    for s in lines:
        if align=='center':c.drawCentredString(x+w/2,yy,s)
        else:c.drawString(x,yy,s)
        yy-=leading
    sh=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h)); tf=sh.text_frame;tf.clear();tf.word_wrap=False
    tf.margin_left=Pt(0);tf.margin_right=Pt(0);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,s in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=s;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=PP_ALIGN.CENTER if align=='center' else PP_ALIGN.LEFT;p.space_before=Pt(0);p.space_after=Pt(0)
    return sh

def box(id,x,y,w,h,label,fill,fg=ink,size=12,bold=False):
    c.setFillColor(HexColor(fill));c.roundRect(x,H-y-h,w,h,6,fill=1,stroke=0)
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sh.name=id;sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background()
    sh.adjustments[0]=.08
    text(x+10,y,w-20,h,label,size,fg,bold,'center');nodes.append(dict(id=id,label=label.replace('\n',' '),x=x,y=y,w=w,h=h))
def seg(x1,y1,x2,y2):
    c.setStrokeColor(HexColor(line));c.setLineWidth(1.1);c.line(x1,H-y1,x2,H-y2)
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2));s.line.color.rgb=rgb(line);s.line.width=Pt(1.1)
text(42,24,700,32,'ORGANIGRAMME',25,navy,True)
text(42,59,900,22,'Structure institutionnelle • Rattachements hiérarchiques',11,'#597080')
# Connectors first, behind the blocks.
seg(595,171,595,277);seg(310,217,880,217)
centers=[175,455,735,1015]
seg(175,277,1015,277)
for x in centers:seg(x,277,x,294)
for j,count in enumerate([5,5,6,3]):
    x=42+280*j;seg(x+133,352,x+133,370);seg(x+133,370,x+4,370);seg(x+4,370,x+4,393+56*(count-1))
    for k in range(count):seg(x+4,393+56*k,x+18,393+56*k)
box('DG',419,99,352,72,'',navy)
portrait=ROOT/'assets/portrait_source.jpeg'
c.drawImage(str(portrait),429,H-161,55,53,mask='auto');slide.shapes.add_picture(str(portrait),Pt(429),Pt(108),width=Pt(55),height=Pt(53))
text(498,108,254,24,'Directeur Executif',17,'#FFFFFF',True,'center')
text(498,136,254,23,'01',12,'#FFFFFF',True,'center')
box('IT',145,195,310,44,'IT',gray,size=13,bold=True)
box('SEC',735,195,310,44,'Secretariat de Direction',gray,size=13,bold=True)
groups=[('Direction Programme et\nPlanification',['Santé Communautaire','SRMNEA/Nutrition &\nMaladies infectieuses','MNT','Santé Climat et environnement','Economie de la santé']),('Direction SERA',['Evaluations & Recherche','Monitorage','Gestion & Analyse des\ndonnées','Publications','Renforcement des Capacités']),('Direction Administration et\nFinance',['Finance','Comptabilité','Trésorerie','Ressources Humaines','Achat /Logistique','Chauffeur']),('Direction Relations Extérieures',['Partenariats','Mob ressources','Communication & Marketing'])]
for j,(d,units) in enumerate(groups):
    x=42+j*280;did=f'D{j+1}';box(did,x,294,266,58,d,blue,'#FFFFFF',12,True);edges.append(('DG',did))
    for k,u in enumerate(units):
        uid=f'{did}-U{k+1}';box(uid,x+18,371+56*k,248,44,u,pale,size=10.5,bold=True);edges.append((did,uid))
edges.extend([('DG','IT'),('DG','SEC')])
# actual chief label retained in data
nodes[0]['label']='Directeur Executif'
seg(42,753,1148,753)
text(42,762,1100,20,'Trait plein : rattachement hiérarchique  •  Les entités de chaque colonne relèvent directement de leur direction.',9,'#597080')
text(42,789,1100,20,'Source : Organigramme vf.pdf, p. 1  |  Lecture des rattachements confirmée par le demandeur  |  30 octobre 2026',8,'#597080')
c.showPage();c.save();prs.save(OUT/'Organigramme_modifiable.pptx')
pdf=fitz.open(OUT/'Organigramme_modernise_A3.pdf');pdf[0].get_pixmap(dpi=300).save(OUT/'Organigramme_haute_resolution.png')
# Verification data: one line per entity and per relation.
rows=[]
for n in nodes:rows.append(['Entité',n['label'],n['label'],'Conforme','Intitulé conservé ; retours à la ligne seuls adaptés.'])
lookup={n['id']:n['label'] for n in nodes}
for p,ch in edges:
    note='Rattachement direct confirmé par le demandeur ; tracé source en chaîne.' if '-U' in ch else 'Rattachement reproduit du document source.'
    rows.append(['Rattachement',lookup[p]+' → '+lookup[ch],lookup[p]+' → '+lookup[ch],'Conforme après clarification' if '-U' in ch else 'Conforme',note])
rows.extend([['Responsables et titulaires','Aucun nom indiqué','Aucun nom ajouté','Conforme','Le portrait ne permet pas d’attribuer un nom.'],['Portrait et repère','Portrait du Directeur Executif ; 01','Portrait extrait ; 01 conservé','Conforme','Conservation expressément demandée.'],['Relations fonctionnelles','Aucune relation explicitement qualifiée','Aucun lien fonctionnel inventé','Non attestées','Absence de pointillés ; ne préjuge pas de l’organisation réelle.'],['Gouvernance / contrôle / antennes','Aucun autre organe explicitement indiqué','Aucun ajout','Non attestés','Aucun développement inventé des sigles IT, SERA, MNT ou SRMNEA.']])
with open(OUT/'Tableau_verification.csv','w',encoding='utf-8-sig',newline='') as f:
    cw=csv.writer(f,delimiter=';');cw.writerow(['Élément','Source / lecture confirmée','Nouvelle version','Statut','Observation']);cw.writerows(rows)
styles=getSampleStyleSheet();styles['Normal'].fontName='Institution';styles['Normal'].fontSize=9;styles['Normal'].leading=13
story=[Paragraph('Analyse et tableau de vérification',styles['Title']),Spacer(1,12)]
intro=['Source examinée : Organigramme vf.pdf, une page A4 paysage. Le livrable principal est vectoriel au format A3 paysage.', 'Inventaire : 26 entités, dont un Directeur Executif, deux fonctions de soutien, quatre directions et dix-neuf entités rattachées à celles-ci. Aucun nom de titulaire, organe supplémentaire de gouvernance, organe de contrôle ou antenne régionale n’est explicitement indiqué.', 'Clarification obtenue : les blocs sous chaque direction sont des entités au même niveau, directement rattachées à leur direction. Les traits en chaîne de la source ne doivent pas être interprétés comme des subordinations successives. Le portrait et le repère 01 sont conservés à la demande du demandeur.', 'Choix graphique : disposition hybride à quatre colonnes, direction exécutive centrée et fonctions de soutien latérales. Les ramifications latérales explicitent les rattachements directs. Une seconde disposition n’a pas été retenue : elle n’apporterait pas de gain de lisibilité sur cette structure.', 'Fidélité : orthographe et sigles officiels conservés, notamment Directeur Executif, Secretariat de Direction, Mob ressources et Achat /Logistique. Aucun intitulé développé, poste ajouté ou nom déduit du portrait. La nature fonctionnelle ou consultative d’un lien n’étant pas attestée, aucun lien pointillé n’est créé.', 'Lecture du tableau : pour les dix-neuf rattachements aux directions, la comparaison repose sur la lecture confirmée par le demandeur et non sur une interprétation non vérifiée du tracé initial.']
for t in intro:story.extend([Paragraph(t,styles['Normal']),Spacer(1,8)])
head=['Élément','Source / lecture confirmée','Nouvelle version','Statut / observation']
data=[[Paragraph(t,styles['Normal']) for t in head]]
for r in rows:data.append([Paragraph(t.replace('&','&amp;'),styles['Normal']) for t in [r[0],r[1],r[2],r[3]+' — '+r[4]]])
t=Table(data,colWidths=[65,150,150,160],repeatRows=1);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EDF5FA')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.3,colors.HexColor('#D8E2EA')),('LEFTPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(t)
SimpleDocTemplate(str(OUT/'Analyse_et_verification.pdf'),pagesize=(595.28,841.89),rightMargin=35,leftMargin=35,topMargin=35,bottomMargin=35).build(story)
(OUT/'structure_verifiee.json').write_text(json.dumps({'nodes':nodes,'edges':edges,'clarification':'Entités de même niveau sous chaque direction ; portrait et 01 conservés.'},ensure_ascii=False,indent=2))
print('Générés :',*[p.name for p in OUT.iterdir()],sep='\n')
